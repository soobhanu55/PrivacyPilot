import asyncio
import json

import httpx
import pytest

from app.rag.corpus import load_regulation_chunks
from app.rag.hybrid_retriever import HybridRetriever
from app.services import qa_service
from app.services.document_reader import read_document


@pytest.fixture(scope="module")
def retriever():
    r = HybridRetriever()  # BM25 only: no model needed
    r.index(load_regulation_chunks())
    return r


def run(coro):
    return asyncio.run(coro)


# ---- Q&A -----------------------------------------------------------------------

def test_ask_returns_one_entry_per_article_without_calling_any_api(retriever):
    def no_network(request):
        raise AssertionError("no API key configured, nothing should be sent")

    client = httpx.AsyncClient(transport=httpx.MockTransport(no_network))
    res = run(qa_service.ask("incident reporting within 24 hours", top_k=3, retriever=retriever, client=client))
    keys = [(c["regulation"], c["article"]) for c in res["clauses"]]
    assert keys[0] == ("NIS2", 23) and len(keys) == len(set(keys)) == 3
    assert res["answer"] is None and res["retrieval"] == "bm25"


def test_ask_can_be_restricted_to_one_regulation(retriever):
    res = run(qa_service.ask("high-risk AI human oversight", ["NIS2"], 5, retriever=retriever))
    assert res["clauses"] and {c["regulation"] for c in res["clauses"]} == {"NIS2"}


def test_ask_sends_cited_excerpts_to_groq_and_returns_the_answer(retriever):
    seen = {}

    def handler(request):
        seen["auth"] = request.headers["Authorization"]
        seen["body"] = json.loads(request.content)
        return httpx.Response(200, json={"choices": [{"message": {"content": "24 hours [NIS2 Art. 23]"}}]})

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    res = run(qa_service.ask("incident reporting within 24 hours", top_k=2, retriever=retriever,
                             groq_key="test-key", client=client))
    assert res["answer"] == "24 hours [NIS2 Art. 23]" and res["generation_error"] is None
    assert seen["auth"] == "Bearer test-key" and "NIS2 Art. 23" in seen["body"]["messages"][1]["content"]
    assert seen["body"]["temperature"] == 0


def test_a_failed_generation_still_returns_the_retrieved_articles(retriever):
    client = httpx.AsyncClient(transport=httpx.MockTransport(lambda r: httpx.Response(500)))
    res = run(qa_service.ask("incident reporting", top_k=2, retriever=retriever, groq_key="k", client=client))
    assert res["clauses"] and res["answer"] is None and res["generation_error"] == "HTTPStatusError"


# ---- document reader -------------------------------------------------------------

def test_reads_text_and_markdown(tmp_path):
    (tmp_path / "a.txt").write_text("Zeile eins\n\nZeile zwei", encoding="utf-8")
    (tmp_path / "b.md").write_text("# Titel", encoding="utf-8")
    assert read_document(tmp_path / "a.txt") == "Zeile eins\n\nZeile zwei"
    assert read_document(tmp_path / "b.md") == "# Titel"


def test_reads_docx_paragraphs(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("Erster Absatz")
    d.add_paragraph("Zweiter Absatz")
    d.save(tmp_path / "policy.docx")
    assert read_document(tmp_path / "policy.docx") == "Erster Absatz\n\nZweiter Absatz"


def test_pdf_without_text_reads_as_empty_not_an_error(tmp_path):
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    with open(tmp_path / "blank.pdf", "wb") as f:
        w.write(f)
    assert read_document(tmp_path / "blank.pdf").strip() == ""


def test_unsupported_type_is_rejected(tmp_path):
    (tmp_path / "x.exe").write_bytes(b"MZ")
    with pytest.raises(ValueError, match="unsupported"):
        read_document(tmp_path / "x.exe")
